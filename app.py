import io
import re
import streamlit as st
import pandas as pd
import plotly.express as px
from difflib import get_close_matches
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# ======================
# CONFIG (harus paling atas)
# ======================
st.set_page_config(page_title="Si Alpha Dashboard", layout="wide")

# ======================
# LOGIN
# ======================
PASSWORD = st.secrets["password"]

if "login" not in st.session_state:
    st.session_state["login"] = False

if not st.session_state["login"]:
    pwd = st.text_input("Masukkan password", type="password")

    if pwd == PASSWORD:
        st.session_state["login"] = True
        st.rerun()
    else:
        st.stop()

# ======================
# HEADER
# ======================
st.markdown("""
<div style='text-align:center'>
<div style='font-size:14px;color:gray;letter-spacing:2px'>DASHBOARD</div>
<div style='font-size:28px;font-weight:600;'>SI ALPHA</div>
</div>
<hr>
""", unsafe_allow_html=True)

# ======================
# LOAD DATA (di-cache 5 menit supaya tidak download ulang tiap klik)
# ======================
url = "https://docs.google.com/spreadsheets/d/1EhwFtO0nBm4w10yZYr18Jft77lVvUtUN/export?format=xlsx"

@st.cache_data(ttl=300)
def load_data(u):
    return pd.read_excel(u)

df = load_data(url).copy()
df.columns = df.columns.str.strip().str.lower()

# ======================
# PROCESS
# ======================
df["tanggal"] = pd.to_datetime(df["tanggal"], errors='coerce')
df["bulan"] = df["tanggal"].dt.to_period("M").astype(str)
df["periode"] = df["tanggal"].dt.strftime("%d-%m-%Y")

df["persentase_perubahan"] = pd.to_numeric(df["persentase_perubahan"], errors='coerce').fillna(0)
df["harga sekarang"] = pd.to_numeric(df["harga sekarang"], errors='coerce')
df["harga sebelum"] = pd.to_numeric(df["harga sebelum"], errors='coerce')
df["catatan"] = df["catatan"].fillna("")
df["catatan"] = df["catatan"].astype(str).apply(
    lambda x: "" if x.strip().lower() in ["tidak ada keterangan", "nan", "-"] else x
)
df["responden"] = df["responden"].fillna("tidak diketahui")

# ======================
# FILTER 1 — DATA UTAMA
# ======================
st.subheader("🔎 Informasi Umum")

df_main = df.copy()

c1,c2,c3,c4 = st.columns(4)

kuesioner_list = ["All"] + sorted(df_main["jenis_kuesioner"].astype(str).unique())
f1_k = c1.selectbox("Kuesioner", kuesioner_list)

if f1_k != "All":
    df_main = df_main[df_main["jenis_kuesioner"] == f1_k]

periode_list = ["All"] + sorted(df_main["periode"].dropna().unique())
f1_b = c2.selectbox("Periode", periode_list)

if f1_b != "All":
    df_main = df_main[df_main["periode"] == f1_b]

komoditas_list = ["All"] + sorted(df_main["komoditas"].astype(str).unique())
f1_ko = c3.selectbox("Komoditas", komoditas_list)

if f1_ko != "All":
    df_main = df_main[df_main["komoditas"] == f1_ko]

kualitas_list = ["All"] + sorted(df_main["kualitas"].astype(str).unique())
f1_ku = c4.selectbox("Kualitas", kualitas_list)

if f1_ku != "All":
    df_main = df_main[df_main["kualitas"] == f1_ku]

# ======================
# TABEL UTAMA
# ======================
st.subheader("📊 Tabel Informasi Umum")

df_main_display = df_main[[
    "tanggal","responden","komoditas","kualitas",
    "harga sekarang","harga sebelum","persentase_perubahan","catatan"
]].copy()

df_main_display["tanggal"] = df_main_display["tanggal"].dt.date
df_main_display["harga sekarang"] = df_main_display["harga sekarang"].map(lambda x: f"Rp {x:,.0f}")
df_main_display["harga sebelum"] = df_main_display["harga sebelum"].map(lambda x: f"Rp {x:,.0f}")
df_main_display["persentase_perubahan"] = df_main_display["persentase_perubahan"].map(lambda x: f"{x:.2f}%")

st.dataframe(df_main_display, use_container_width=True)

# ======================
# FILTER 2 — ANALISIS
# ======================
st.subheader("📝 Analisis")

df_analysis_filter = df.copy()

a1,a2,a3,a4 = st.columns(4)

fa_k = a1.selectbox("Kuesioner", ["All"] + sorted(df["jenis_kuesioner"].astype(str).unique()), key="a1")

if fa_k != "All":
    df_analysis_filter = df_analysis_filter[df_analysis_filter["jenis_kuesioner"] == fa_k]

periode_dyn = ["All"] + sorted(df_analysis_filter["periode"].dropna().unique())
fa_b = a2.selectbox("Periode", periode_dyn, key="a2")

if fa_b != "All":
    df_analysis_filter = df_analysis_filter[df_analysis_filter["periode"] == fa_b]

kom_dyn = ["All"] + sorted(df_analysis_filter["komoditas"].astype(str).unique())
fa_ko = a3.selectbox("Komoditas", kom_dyn, key="a3")

if fa_ko != "All":
    df_analysis_filter = df_analysis_filter[df_analysis_filter["komoditas"] == fa_ko]

kualitas_dyn = ["All"] + sorted(df_analysis_filter["kualitas"].astype(str).unique())
fa_ku = a4.selectbox("Kualitas", kualitas_dyn, key="a4")

if fa_ku != "All":
    df_analysis_filter = df_analysis_filter[df_analysis_filter["kualitas"] == fa_ku]

# ======================
# ANALISIS DATA
# ======================
df_analysis = df_analysis_filter[[
    "komoditas","kualitas","responden","persentase_perubahan","catatan"
]].copy()

# ======================
# CLEAN TEXT
# ======================
def clean_text(text):
    text = text.lower().strip()
    text = re.sub(r'\b(\w+)( \1\b)+', r'\1', text)
    text = re.sub(r'\s+', ' ', text)
    return text

# ======================
# INSIGHT
# ======================
st.subheader("🧠 Insight")

if not df_analysis.empty:

    rata2 = df_analysis["persentase_perubahan"].mean().round(2)

    df_analysis["abs"] = df_analysis["persentase_perubahan"].abs()
    df_analysis = df_analysis.sort_values("abs", ascending=False)

    top = df_analysis.iloc[0]

    arah = "inflasi" if rata2 > 0 else "deflasi"

    df_top = df_analysis[
        (df_analysis["komoditas"] == top["komoditas"]) &
        (df_analysis["kualitas"] == top["kualitas"])
    ]

    catatan_list = [c for c in df_top["catatan"].dropna().tolist() if c.strip()]
    cleaned = [clean_text(c) for c in catatan_list]

    unik = []
    for kalimat in cleaned:
        if not any(kalimat in u or u in kalimat for u in unik):
            unik.append(kalimat)

    sebab = ", ".join(unik)
    sebab = sebab.capitalize() if sebab else "Tidak ada keterangan utama"

    narasi = (
        f"Pada {fa_b}, terjadi {arah} sebesar {rata2:.2f}%. "
        f"Komoditas utama: {top['komoditas']} ({top['kualitas']}) "
        f"dengan perubahan {top['persentase_perubahan']:.2f}%. "
        f"Penyebab: {sebab}."
    )

    st.info(narasi)

# ======================
# TABEL INFLASI DEFLASI
# ======================
df_naik = df_analysis[df_analysis["persentase_perubahan"] > 0].copy()
df_turun = df_analysis[df_analysis["persentase_perubahan"] < 0].copy()

df_naik["persentase_perubahan"] = df_naik["persentase_perubahan"].map(lambda x: f"{x:.2f}%")
df_turun["persentase_perubahan"] = df_turun["persentase_perubahan"].map(lambda x: f"{x:.2f}%")

c1,c2 = st.columns(2)

with c1:
    st.markdown("#### 🔴 Inflasi")
    st.dataframe(df_naik[["komoditas","kualitas","responden","persentase_perubahan","catatan"]], use_container_width=True)

with c2:
    st.markdown("#### 🟢 Deflasi")
    st.dataframe(df_turun[["komoditas","kualitas","responden","persentase_perubahan","catatan"]], use_container_width=True)

# ======================
# HARGA TIDUR
# ======================
df_tidur = df[df["persentase_perubahan"] == 0].copy()
df_tidur["bulan_dt"] = df_tidur["tanggal"].dt.to_period("M")

tidur_group = df_tidur.groupby(["komoditas","kualitas"]).agg(
    bulan_dt=("bulan_dt","nunique"),
    responden=("responden", lambda x: ", ".join(sorted(set(x))))
).reset_index()
tidur_final = tidur_group[tidur_group["bulan_dt"] >= 3]

st.subheader("🛌 Harga Tidur")
st.dataframe(tidur_final[["komoditas","kualitas","responden"]], use_container_width=True)

# ======================
# GRAFIK
# ======================
st.subheader("📈 Tren Harga")

kom = st.selectbox("Pilih Komoditas", sorted(df["komoditas"].unique()))

df_grafik = df[df["komoditas"] == kom]
df_grafik = df_grafik.groupby(["tanggal","kualitas"], as_index=False)["harga sekarang"].mean()

fig = px.line(df_grafik, x="tanggal", y="harga sekarang", color="kualitas")

st.plotly_chart(fig, use_container_width=True)

# ======================
# CEK ANGKA SEMENTARA (FITUR BARU)
# ======================
st.subheader("✅ Cek Angka Sementara vs Catatan")

# kolom dari file angka sementara -> judul kolom di tabel
KOLOM_TAMPIL = {
    "nama": "Nama",
    "flag": "Flag",
    "ihk": "IHK",
    "inf(mom)": "Inf (MoM)",
    "inf(ytd)": "Inf (YTD)",
    "inf(yoy)": "Inf (YoY)",
    "andil(mom)": "Andil (MoM)",
    "andil(ytd)": "Andil (YTD)",
    "andil(yoy)": "Andil (YoY)",
}

def norm(x):
    return re.sub(r"\s+", " ", str(x).strip().lower())

def baca_sementara(file):
    raw = pd.read_excel(file, header=None, dtype=str)
    # cari baris header (yang berisi kolom "Nama")
    baris = raw.apply(lambda r: r.astype(str).str.strip().str.lower().eq("nama").any(), axis=1)
    if not baris.any():
        return None
    hdr = baris.idxmax()
    d = pd.read_excel(file, skiprows=hdr, dtype={"Bulan": str, "Tahun": str, "Kode": str})
    d.columns = d.columns.astype(str).str.strip().str.lower()
    return d

def catatan_responden(g):
    """Gabungkan catatan satu responden (bisa beberapa kualitas/minggu)."""
    bagian = []
    for _, x in g.sort_values(["kualitas", "tanggal"]).iterrows():
        cat = x["catatan"].strip() if isinstance(x["catatan"], str) else ""
        if cat:
            bagian.append(f"[{x['kualitas']}, {x['persentase_perubahan']:+.2f}%] {cat}")
    bagian = list(dict.fromkeys(bagian))  # buang duplikat, urutan tetap
    return " | ".join(bagian) if bagian else "Tidak ada catatan"

file_sementara = st.file_uploader(
    "Upload Excel angka sementara", type=["xlsx", "xls"], key="upload_sementara"
)

if file_sementara is not None:
    df_sem = baca_sementara(file_sementara)

    kolom_wajib = ["nama", "tahun", "bulan"]
    if df_sem is None or any(k not in df_sem.columns for k in kolom_wajib):
        st.error("Format file tidak dikenali. Butuh minimal kolom: Nama, Tahun, Bulan.")
    else:
        # kolom angka jadi numerik
        for k in KOLOM_TAMPIL:
            if k not in ("nama", "flag") and k in df_sem.columns:
                df_sem[k] = pd.to_numeric(df_sem[k], errors="coerce")

        df_sem = df_sem[df_sem["nama"].notna()].copy()
        df_sem["_kom"] = df_sem["nama"].apply(norm)

        # bulan otomatis dari file (format sama dengan kolom 'bulan' di df: 2026-09)
        bulan_file = f"{str(df_sem['tahun'].iloc[0]).strip()}-{str(df_sem['bulan'].iloc[0]).strip().zfill(2)}"
        bulan_list = sorted(df["bulan"].dropna().unique())
        idx = bulan_list.index(bulan_file) if bulan_file in bulan_list else len(bulan_list) - 1

        c1, c2 = st.columns(2)
        bulan_pilih = c1.selectbox("Bulan data entri", bulan_list, index=idx, key="bulan_sem")
        ambang = c2.number_input(
            "Tampilkan hanya komoditas dengan |Inf (MoM)| ≥ (%) — isi 0 untuk semua",
            min_value=0.0, value=0.0, step=0.5, key="ambang_sem"
        )

        if bulan_file not in bulan_list:
            st.warning(f"Bulan di file ({bulan_file}) belum ada di data entri. Menampilkan {bulan_pilih}.")

        if ambang > 0 and "inf(mom)" in df_sem.columns:
            df_sem = df_sem[df_sem["inf(mom)"].abs() >= ambang].copy()

        # data entri bulan tsb
        df["_kom"] = df["komoditas"].apply(norm)
        df_bulan = df[df["bulan"] == bulan_pilih]
        nama_entri = df_bulan["_kom"].unique().tolist()

        # matching: persis dulu, lalu mirip
        def cocokkan(k):
            if k in nama_entri:
                return k
            m = get_close_matches(k, nama_entri, n=1, cutoff=0.85)
            return m[0] if m else None

        df_sem["_match"] = df_sem["_kom"].apply(cocokkan)

        # ---------- bangun tabel: 1 baris per responden ----------
        rows = []
        for _, r in df_sem.iterrows():
            base = {label: r[k] for k, label in KOLOM_TAMPIL.items() if k in df_sem.columns}
            d = df_bulan[df_bulan["_kom"] == r["_match"]] if r["_match"] else df_bulan.iloc[0:0]

            if d.empty:
                rows.append({**base, "Responden": "-", "Catatan": "Tidak ada data di entri"})
            else:
                for responden, g in d.groupby("responden"):
                    rows.append({**base, "Responden": responden, "Catatan": catatan_responden(g)})

        df_out = pd.DataFrame(rows)

        # ---------- ringkasan ----------
        n_total = df_out["Nama"].nunique()
        n_ada = df_out.loc[df_out["Catatan"] != "Tidak ada data di entri", "Nama"].nunique()
        n_cat = df_out.loc[~df_out["Catatan"].isin(["Tidak ada catatan", "Tidak ada data di entri"]), "Nama"].nunique()

        m1, m2, m3 = st.columns(3)
        m1.metric("Komoditas di file", n_total)
        m2.metric("Ada di data entri", n_ada)
        m3.metric("Yang punya catatan", n_cat)

        f1, f2, f3 = st.columns([2, 2, 2])
        cari = f1.text_input("Cari komoditas", key="cari_sem")
        tampil_kosong = f2.checkbox("Tampilkan juga komoditas yang tidak ada di entri", key="kosong_sem")
        mode = f3.radio("Tampilan", ["Tabel interaktif", "Tabel teks penuh"],
                        horizontal=True, key="mode_sem")

        df_show = df_out.copy()
        if not tampil_kosong:
            df_show = df_show[df_show["Catatan"] != "Tidak ada data di entri"]
        if cari.strip():
            df_show = df_show[df_show["Nama"].str.contains(cari.strip(), case=False, na=False)]

        # ---------- tampil ----------
        if df_show.empty:
            st.info("Tidak ada data yang sesuai.")
        else:
            if mode == "Tabel interaktif":
                st.dataframe(
                    df_show, use_container_width=True, hide_index=True,
                    column_config={"Catatan": st.column_config.TextColumn("Catatan", width="large")}
                )
                st.caption("Teks catatan panjang bisa terpotong. Pilih 'Tabel teks penuh' untuk melihat semuanya.")
            else:
                st.table(df_show.reset_index(drop=True))

            # ---------- download Excel ----------
            buf = io.BytesIO()
            with pd.ExcelWriter(buf, engine="openpyxl") as w:
                df_show.to_excel(w, index=False, sheet_name="Angka Sementara")
                ws = w.sheets["Angka Sementara"]
                for i, nama_kol in enumerate(df_show.columns, start=1):
                    if nama_kol == "Catatan":
                        lebar = 90
                    elif nama_kol == "Nama":
                        lebar = 35
                    elif nama_kol == "Responden":
                        lebar = 22
                    else:
                        lebar = 13
                    ws.column_dimensions[get_column_letter(i)].width = lebar
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")

            st.download_button(
                "⬇️ Download Excel",
                buf.getvalue(),
                file_name=f"angka_sementara_catatan_{bulan_pilih}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
